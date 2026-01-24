"""
Report Generation Module

Generates formatted reports in multiple formats:
- HTML for web display
- Excel/CSV for data export
- PDF for professional documentation
"""

from typing import Dict, List, Optional
from datetime import datetime
import io
import os

# For Excel export
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

# For PDF export
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December"
]


class ReportGenerator:
    """Generate weather day analysis reports in multiple formats."""

    def __init__(
        self,
        station_info: Dict,
        location_info: Dict,
        analysis_data: Dict,
        period_name: str
    ):
        """
        Initialize report generator.

        Args:
            station_info: NOAA station information
            location_info: Geocoded location information
            analysis_data: Processed precipitation analysis data
            period_name: Name of the analysis period (e.g., "5-Year")
        """
        self.station = station_info
        self.location = location_info
        self.data = analysis_data
        self.period_name = period_name
        self.generated_at = datetime.now()

    def _format_value(self, value: Optional[int], is_average: bool = False, is_approximate: bool = False) -> str:
        """Format a data value for display."""
        if value is None:
            return "-"

        if is_average:
            if is_approximate:
                return f"~{value}"
            return str(value)

        return str(value)

    def _get_header_text(self) -> Dict[str, str]:
        """Generate header text for the report."""
        station_name = self.station.get("name", "Unknown Station")
        distance = self.station.get("distance_miles", "?")
        direction = self.station.get("direction", "")
        location = self.location.get("full_location", self.location.get("zip_code", ""))

        # Year range
        years = self.data.get("years", [])
        year_range = f"{min(years)}-{max(years)}" if years else "Unknown"

        # Data quality notes
        quality = self.data.get("data_quality", {})
        missing_years = quality.get("missing_years", [])
        partial_years = quality.get("partial_years", [])

        quality_note = ""
        if missing_years:
            quality_note = f"Station is lacking data for {', '.join(map(str, missing_years))}; adjusted analysis accordingly."
        elif partial_years:
            quality_note = f"Station has partial data for {', '.join(map(str, partial_years))}; some months may be incomplete."

        return {
            "station_name": station_name,
            "distance": distance,
            "direction": direction,
            "location": location,
            "year_range": year_range,
            "quality_note": quality_note,
            "period_name": self.period_name
        }

    def generate_html(self) -> str:
        """
        Generate HTML report.

        Returns:
            HTML string
        """
        header = self._get_header_text()
        years = self.data.get("years", [])
        monthly_data = self.data.get("monthly_rain_days", {})
        averages = self.data.get("monthly_averages", {})
        yearly_totals = self.data.get("yearly_totals", {})
        avg_total = self.data.get("average_total", (None, False))

        # Build table rows
        rows_html = ""
        for month_idx, month_name in enumerate(MONTH_NAMES, 1):
            row = f"<tr><td class='month'>{month_name}</td>"

            for year in years:
                value = monthly_data.get(year, {}).get(month_idx)
                formatted = self._format_value(value)
                css_class = "missing" if value is None else ""
                row += f"<td class='{css_class}'>{formatted}</td>"

            # Average column
            avg_value, is_approx = averages.get(month_idx, (None, False))
            avg_formatted = self._format_value(avg_value, is_average=True, is_approximate=is_approx)
            row += f"<td class='average'>{avg_formatted}</td></tr>\n"

            rows_html += row

        # Total row
        total_row = "<tr class='total-row'><td class='month'><strong>TOTAL</strong></td>"
        for year in years:
            value = yearly_totals.get(year)
            formatted = self._format_value(value)
            css_class = "missing" if value is None else ""
            total_row += f"<td class='{css_class}'><strong>{formatted}</strong></td>"

        avg_total_val, avg_total_approx = avg_total
        avg_total_formatted = self._format_value(avg_total_val, is_average=True, is_approximate=avg_total_approx)
        total_row += f"<td class='average'><strong>{avg_total_formatted}</strong></td></tr>"

        # Year headers
        year_headers = "".join(f"<th>{year}</th>" for year in years)

        html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>Weather Day Analysis Report</title>
    <style>
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px;
            background-color: #f5f5f5;
        }}
        .report-container {{
            background: white;
            padding: 30px;
            border-radius: 8px;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }}
        .header {{
            margin-bottom: 25px;
            padding-bottom: 15px;
            border-bottom: 2px solid #2c5282;
        }}
        .header h1 {{
            color: #2c5282;
            margin: 0 0 10px 0;
            font-size: 1.5em;
        }}
        .station-info {{
            color: #4a5568;
            font-size: 1.1em;
            margin-bottom: 10px;
        }}
        .quality-note {{
            background-color: #fef3c7;
            border-left: 4px solid #f59e0b;
            padding: 10px 15px;
            margin: 15px 0;
            color: #92400e;
        }}
        .data-table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }}
        .data-table th, .data-table td {{
            padding: 10px 12px;
            text-align: center;
            border: 1px solid #e2e8f0;
        }}
        .data-table th {{
            background-color: #2c5282;
            color: white;
            font-weight: 600;
        }}
        .data-table th:first-child {{
            text-align: left;
        }}
        .data-table .month {{
            text-align: left;
            font-weight: 500;
            background-color: #f7fafc;
        }}
        .data-table .average {{
            background-color: #ebf4ff;
            font-weight: 600;
        }}
        .data-table .missing {{
            color: #a0aec0;
        }}
        .data-table .total-row {{
            background-color: #edf2f7;
        }}
        .data-table tr:hover {{
            background-color: #f0f4f8;
        }}
        .footer {{
            margin-top: 25px;
            padding-top: 15px;
            border-top: 1px solid #e2e8f0;
            font-size: 0.9em;
            color: #718096;
        }}
        .legend {{
            margin-top: 15px;
            font-size: 0.9em;
            color: #4a5568;
        }}
        .legend span {{
            margin-right: 20px;
        }}
        @media print {{
            body {{
                background: white;
            }}
            .report-container {{
                box-shadow: none;
            }}
        }}
    </style>
</head>
<body>
    <div class="report-container">
        <div class="header">
            <h1>Historical Rain Days with &ge;0.5" Precipitation</h1>
            <div class="station-info">
                The nearest NOAA station is <strong>{header['station_name']}</strong>,
                approximately <strong>{header['distance']} miles {header['direction']}</strong>
                of {header['location']}.
            </div>
            <div class="station-info">
                {header['station_name']} Weather Station ({header['year_range']}) &mdash; Full Calendar Year Data
            </div>
        </div>

        {"<div class='quality-note'>" + header['quality_note'] + "</div>" if header['quality_note'] else ""}

        <table class="data-table">
            <thead>
                <tr>
                    <th>Month</th>
                    {year_headers}
                    <th>{header['period_name']} Avg</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
                {total_row}
            </tbody>
        </table>

        <div class="legend">
            <span><strong>~</strong> = Approximate (some years missing data)</span>
            <span><strong>-</strong> = No data available</span>
        </div>

        <div class="footer">
            <p>Data source: NOAA Climate Data Online (CDO) - Global Historical Climatology Network Daily (GHCND)</p>
            <p>Report generated: {self.generated_at.strftime('%B %d, %Y at %I:%M %p')}</p>
            <p>Precipitation threshold: &ge;0.5 inches per day</p>
        </div>
    </div>
</body>
</html>
"""
        return html

    def generate_excel(self) -> bytes:
        """
        Generate Excel report.

        Returns:
            Excel file as bytes
        """
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = "Rain Day Analysis"

        header = self._get_header_text()
        years = self.data.get("years", [])
        monthly_data = self.data.get("monthly_rain_days", {})
        averages = self.data.get("monthly_averages", {})
        yearly_totals = self.data.get("yearly_totals", {})
        avg_total = self.data.get("average_total", (None, False))

        # Styles
        header_font = Font(bold=True, size=14, color="FFFFFF")
        subheader_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
        avg_fill = PatternFill(start_color="EBF4FF", end_color="EBF4FF", fill_type="solid")
        total_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title section
        ws.merge_cells('A1:' + get_column_letter(len(years) + 2) + '1')
        ws['A1'] = "Historical Rain Days with ≥0.5\" Precipitation"
        ws['A1'].font = Font(bold=True, size=16)

        ws.merge_cells('A2:' + get_column_letter(len(years) + 2) + '2')
        ws['A2'] = f"Station: {header['station_name']} ({header['distance']} miles {header['direction']} of {header['location']})"
        ws['A2'].font = Font(size=11)

        ws.merge_cells('A3:' + get_column_letter(len(years) + 2) + '3')
        ws['A3'] = f"Period: {header['year_range']} - Full Calendar Year Data"
        ws['A3'].font = Font(size=11)

        if header['quality_note']:
            ws.merge_cells('A4:' + get_column_letter(len(years) + 2) + '4')
            ws['A4'] = header['quality_note']
            ws['A4'].font = Font(size=10, italic=True)
            data_start_row = 6
        else:
            data_start_row = 5

        # Column headers
        row = data_start_row
        ws.cell(row=row, column=1, value="Month").font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.cell(row=row, column=1).border = thin_border

        for col, year in enumerate(years, 2):
            cell = ws.cell(row=row, column=col, value=year)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        avg_col = len(years) + 2
        ws.cell(row=row, column=avg_col, value=f"{header['period_name']} Avg").font = header_font
        ws.cell(row=row, column=avg_col).fill = header_fill
        ws.cell(row=row, column=avg_col).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=avg_col).border = thin_border

        # Data rows
        for month_idx, month_name in enumerate(MONTH_NAMES, 1):
            row += 1
            ws.cell(row=row, column=1, value=month_name).border = thin_border

            for col, year in enumerate(years, 2):
                value = monthly_data.get(year, {}).get(month_idx)
                cell = ws.cell(row=row, column=col, value=value if value is not None else "-")
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border

            avg_value, is_approx = averages.get(month_idx, (None, False))
            display_val = f"~{avg_value}" if is_approx and avg_value is not None else (avg_value if avg_value is not None else "-")
            cell = ws.cell(row=row, column=avg_col, value=display_val)
            cell.fill = avg_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        # Total row
        row += 1
        ws.cell(row=row, column=1, value="TOTAL").font = subheader_font
        ws.cell(row=row, column=1).fill = total_fill
        ws.cell(row=row, column=1).border = thin_border

        for col, year in enumerate(years, 2):
            value = yearly_totals.get(year)
            cell = ws.cell(row=row, column=col, value=value if value is not None else "-")
            cell.font = subheader_font
            cell.fill = total_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        avg_total_val, avg_total_approx = avg_total
        display_val = f"~{avg_total_val}" if avg_total_approx and avg_total_val is not None else (avg_total_val if avg_total_val is not None else "-")
        cell = ws.cell(row=row, column=avg_col, value=display_val)
        cell.font = subheader_font
        cell.fill = avg_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

        # Footer
        row += 2
        ws.cell(row=row, column=1, value="Legend: ~ = Approximate (some years missing data), - = No data available")
        row += 1
        ws.cell(row=row, column=1, value=f"Data source: NOAA Climate Data Online (CDO)")
        row += 1
        ws.cell(row=row, column=1, value=f"Report generated: {self.generated_at.strftime('%B %d, %Y')}")

        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        for col in range(2, len(years) + 3):
            ws.column_dimensions[get_column_letter(col)].width = 10

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def generate_csv(self) -> str:
        """
        Generate CSV report.

        Returns:
            CSV string
        """
        header = self._get_header_text()
        years = self.data.get("years", [])
        monthly_data = self.data.get("monthly_rain_days", {})
        averages = self.data.get("monthly_averages", {})
        yearly_totals = self.data.get("yearly_totals", {})
        avg_total = self.data.get("average_total", (None, False))

        lines = []

        # Header info
        lines.append(f"# Historical Rain Days with >=0.5\" Precipitation")
        lines.append(f"# Station: {header['station_name']} ({header['distance']} miles {header['direction']} of {header['location']})")
        lines.append(f"# Period: {header['year_range']}")
        if header['quality_note']:
            lines.append(f"# Note: {header['quality_note']}")
        lines.append("")

        # Column headers
        headers = ["Month"] + [str(y) for y in years] + [f"{header['period_name']} Avg"]
        lines.append(",".join(headers))

        # Data rows
        for month_idx, month_name in enumerate(MONTH_NAMES, 1):
            row = [month_name]

            for year in years:
                value = monthly_data.get(year, {}).get(month_idx)
                row.append(str(value) if value is not None else "-")

            avg_value, is_approx = averages.get(month_idx, (None, False))
            if avg_value is not None:
                row.append(f"~{avg_value}" if is_approx else str(avg_value))
            else:
                row.append("-")

            lines.append(",".join(row))

        # Total row
        total_row = ["TOTAL"]
        for year in years:
            value = yearly_totals.get(year)
            total_row.append(str(value) if value is not None else "-")

        avg_total_val, avg_total_approx = avg_total
        if avg_total_val is not None:
            total_row.append(f"~{avg_total_val}" if avg_total_approx else str(avg_total_val))
        else:
            total_row.append("-")

        lines.append(",".join(total_row))

        # Footer
        lines.append("")
        lines.append(f"# Generated: {self.generated_at.strftime('%Y-%m-%d')}")
        lines.append(f"# Data source: NOAA Climate Data Online")

        return "\n".join(lines)

    def generate_pdf(self) -> bytes:
        """
        Generate PDF report.

        Returns:
            PDF file as bytes
        """
        if not PDF_AVAILABLE:
            raise ImportError("reportlab is required for PDF export. Install with: pip install reportlab")

        header = self._get_header_text()
        years = self.data.get("years", [])
        monthly_data = self.data.get("monthly_rain_days", {})
        averages = self.data.get("monthly_averages", {})
        yearly_totals = self.data.get("yearly_totals", {})
        avg_total = self.data.get("average_total", (None, False))

        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter) if len(years) > 6 else letter,
            rightMargin=0.5*inch,
            leftMargin=0.5*inch,
            topMargin=0.5*inch,
            bottomMargin=0.5*inch
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=12,
            textColor=colors.HexColor('#2C5282')
        )
        info_style = ParagraphStyle(
            'Info',
            parent=styles['Normal'],
            fontSize=10,
            spaceAfter=6
        )
        note_style = ParagraphStyle(
            'Note',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#92400E'),
            backColor=colors.HexColor('#FEF3C7'),
            borderPadding=5
        )
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.gray
        )

        elements = []

        # Title
        elements.append(Paragraph("Historical Rain Days with &ge;0.5\" Precipitation", title_style))

        # Station info
        elements.append(Paragraph(
            f"The nearest NOAA station is <b>{header['station_name']}</b>, "
            f"approximately <b>{header['distance']} miles {header['direction']}</b> of {header['location']}.",
            info_style
        ))
        elements.append(Paragraph(
            f"{header['station_name']} Weather Station ({header['year_range']}) - Full Calendar Year Data",
            info_style
        ))

        if header['quality_note']:
            elements.append(Spacer(1, 6))
            elements.append(Paragraph(header['quality_note'], note_style))

        elements.append(Spacer(1, 12))

        # Build table data
        table_data = []

        # Header row
        header_row = ["Month"] + [str(y) for y in years] + [f"{header['period_name']} Avg"]
        table_data.append(header_row)

        # Data rows
        for month_idx, month_name in enumerate(MONTH_NAMES, 1):
            row = [month_name]

            for year in years:
                value = monthly_data.get(year, {}).get(month_idx)
                row.append(str(value) if value is not None else "-")

            avg_value, is_approx = averages.get(month_idx, (None, False))
            if avg_value is not None:
                row.append(f"~{avg_value}" if is_approx else str(avg_value))
            else:
                row.append("-")

            table_data.append(row)

        # Total row
        total_row = ["TOTAL"]
        for year in years:
            value = yearly_totals.get(year)
            total_row.append(str(value) if value is not None else "-")

        avg_total_val, avg_total_approx = avg_total
        if avg_total_val is not None:
            total_row.append(f"~{avg_total_val}" if avg_total_approx else str(avg_total_val))
        else:
            total_row.append("-")

        table_data.append(total_row)

        # Create table
        col_widths = [80] + [50] * len(years) + [70]
        table = Table(table_data, colWidths=col_widths)

        # Table styling
        table_style = TableStyle([
            # Header row
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C5282')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),

            # Data rows
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ALIGN', (1, 1), (-1, -1), 'CENTER'),
            ('ALIGN', (0, 1), (0, -1), 'LEFT'),

            # Average column
            ('BACKGROUND', (-1, 1), (-1, -1), colors.HexColor('#EBF4FF')),

            # Total row
            ('BACKGROUND', (0, -1), (-2, -1), colors.HexColor('#EDF2F7')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),

            # Grid
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ])

        table.setStyle(table_style)
        elements.append(table)

        # Legend
        elements.append(Spacer(1, 12))
        elements.append(Paragraph(
            "<b>~</b> = Approximate (some years missing data) &nbsp;&nbsp;&nbsp; "
            "<b>-</b> = No data available",
            footer_style
        ))

        # Footer
        elements.append(Spacer(1, 12))
        elements.append(Paragraph(
            f"Data source: NOAA Climate Data Online (CDO) - Global Historical Climatology Network Daily (GHCND)",
            footer_style
        ))
        elements.append(Paragraph(
            f"Report generated: {self.generated_at.strftime('%B %d, %Y at %I:%M %p')}",
            footer_style
        ))
        elements.append(Paragraph(
            f"Precipitation threshold: &ge;0.5 inches per day",
            footer_style
        ))

        doc.build(elements)
        output.seek(0)
        return output.getvalue()


def generate_report(
    station_info: Dict,
    location_info: Dict,
    analysis_data: Dict,
    period_name: str,
    format: str = "html"
) -> bytes | str:
    """
    Convenience function to generate a report.

    Args:
        station_info: NOAA station information
        location_info: Geocoded location information
        analysis_data: Processed precipitation analysis data
        period_name: Name of the analysis period
        format: Output format ("html", "excel", "csv", or "pdf")

    Returns:
        Report content (string for html/csv, bytes for excel/pdf)
    """
    generator = ReportGenerator(station_info, location_info, analysis_data, period_name)

    if format == "html":
        return generator.generate_html()
    elif format == "excel":
        return generator.generate_excel()
    elif format == "csv":
        return generator.generate_csv()
    elif format == "pdf":
        return generator.generate_pdf()
    else:
        raise ValueError(f"Unsupported format: {format}")
